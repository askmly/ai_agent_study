from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

wb = Workbook()
ws = wb.active
ws.title = "Day12-24压缩版"

headers = [
    "新天", "对应原天", "阶段", "主题",
    "理论30-45min（概念/全局）",
    "上午/下午实践",
    "晚上/产出",
    "完成"
]

rows = [
    ["新Day12", "原Day12+Day13", "Prompt工程", "Prompt工程 + 结构化输出",
     "System/User/Assistant、Few-shot、角色设定、JSON Schema、输出约束",
     "写翻译助手/代码审查/面试官模拟三个Prompt；让模型按JSON输出并解析",
     "对比输出质量；整理 prompts.md；Git commit", ""],

    ["新Day13", "原Day13晚+Day14", "Prompt工程", "Prompt优化 + 第2周整理",
     "LLM应用架构图：前端→API→LLM→输出",
     "优化Prompt；整理API/Prompt代码；推GitHub，写第二周README",
     "复盘：让代码调用AI；第二周仓库+日志", ""],

    ["新Day14", "原Day15", "工具调用", "工具调用原理 + 最小调用",
     "Function Calling如何让LLM连接外部世界",
     "定义Python函数作为工具；用API实现最小工具调用",
     "工具定义示例+最小调用脚本；Git commit", ""],

    ["新Day15", "原Day16", "Agent入门", "天气助手Agent",
     "Agent vs Chatbot；感知-决策-行动循环",
     "让LLM决定何时调用天气函数；跑通“问天气→调用→返回”",
     "处理参数错误；天气助手Agent；Git commit", ""],

    ["新Day16", "原Day17", "RAG", "RAG基础",
     "向量检索、Embedding、相似度、向量数据库",
     "文档加载/切分/向量化/检索；用LangChain加载Markdown/PDF",
     "文档加载脚本；RAG笔记；Git commit", ""],

    ["新Day17", "原Day18", "RAG", "笔记问答助手",
     "RAG进阶：切分策略、召回与重排、上下文注入",
     "搭建向量库；实现基于文档的问答",
     "测试边界问题；RAG问答助手；Git commit", ""],

    ["新Day18", "原Day19+Day21上午", "ReAct", "ReAct + 知识框架",
     "Thought/Action/Observation；推理与行动结合",
     "用LangChain Agent Executor跑最小ReAct；看日志理解决策",
     "ReAct最小示例；知识框架图初稿", ""],

    ["新Day19", "原Day20+Day21下午", "多工具Agent", "多工具Agent + 架构总览",
     "多工具与规划：任务分解、工具选择",
     "配计算器+查资料两个工具；让Agent多步推理选工具",
     "修Bug；Agent架构总览图+问题清单；Git commit", ""],

    ["新Day20", "原Day22+Day23上午", "项目启动", "项目启动+骨架+核心流程",
     "项目结构、模块化、依赖管理、系统集成",
     "选方案A知识库问答/方案B多工具办公助手；建项目结构；实现主流程",
     "项目骨架+可运行最小版本；Git初始化", ""],

    ["新Day21", "原Day23下午+Day24", "项目联调", "联调+修Bug+错误处理",
     "调试、日志、异常、可观测性",
     "联调API/工具/向量库；加try/except、日志；处理异常输入",
     "优化Prompt；稳定版v1；Git commit", ""],

    ["新Day22", "原Day25+Day26", "交付文档", "README+演示+GitHub整理",
     "文档与交付：README、可复现性、版本发布",
     "写项目说明/运行方式/示例；补已知限制；录2-3分钟演示视频",
     "完整README+演示视频+GitHub项目", ""],

    ["新Day23", "原Day27+Day28", "答辩准备", "模拟现场改功能+答辩准备",
     "测试与评估：单元测试、Agent评估指标；安全与成本",
     "自出题30分钟内改功能；练Cursor快速定位；准备答辩三问",
     "模拟考核记录+答辩提纲；Git commit", ""],

    ["新Day24", "原Day29+Day30", "试岗交付", "干净环境跑通+试岗交付",
     "部署与运维：环境变量、Docker基础、云服务概念",
     "删虚拟环境重装依赖；按README完整跑一遍；修复问题",
     "整理练习代码；检查GitHub提交记录/学习日志；试岗交付包", ""],
]

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# 标题
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
ws["A1"] = "AI编程与Agent开发 压缩后13天学习计划（新Day12-Day24）"
ws["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="2F5597")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 32

# 副标题
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
ws["A2"] = "每天：理论30-45min + 上午3h学 + 下午3h练 + 晚上1-2h AI辅助/复盘；每50分钟休息10分钟；每天至少1次Git commit。完成列可下拉选择 ☐ / √。"
ws["A2"].font = Font(name="微软雅黑", size=9, italic=True, color="555555")
ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 30

# 表头
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col, value=h)
    cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="4472C4")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[3].height = 36

# 数据
for r, row_data in enumerate(rows, start=4):
    for c, val in enumerate(row_data, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = Font(name="微软雅黑", size=10)
        cell.border = border
        if c in (1, 2, 8):
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if r % 2 == 0:
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F7FF")
    ws.row_dimensions[r].height = 60

# 完成列下拉
dv = DataValidation(type="list", formula1='"☐,√"', allow_blank=True)
dv.error = "请选择 ☐ 或 √"
dv.errorTitle = "输入错误"
dv.prompt = "选择完成状态"
dv.promptTitle = "完成"
ws.add_data_validation(dv)
dv.add("H4:H16")

# 列宽
widths = [10, 16, 12, 22, 34, 42, 36, 8]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 冻结与筛选
ws.freeze_panes = "A4"
ws.auto_filter.ref = "A3:H16"

# 页面设置
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = 9  # A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.print_title_rows = "1:3"
ws.print_area = "A1:H16"
ws.page_margins.left = 0.3
ws.page_margins.right = 0.3
ws.page_margins.top = 0.4
ws.page_margins.bottom = 0.4
ws.print_options.horizontalCentered = True
ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
ws.oddFooter.center.size = 9

wb.save("修改后_day12-day24学习计划.xlsx")
print("已生成:修改后_day12-day24学习计划.xlsx")